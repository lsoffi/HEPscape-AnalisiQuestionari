import copy
import csv
import json
from pathlib import Path
from types import SimpleNamespace
import pytest
from PIL import Image
from openpyxl import load_workbook
from hepscape.schema import QUESTIONS, OPTIONS, load, validate, coded
from hepscape.workbook import export_workbook, read_workbook
from hepscape.review import export_review, apply_review
from hepscape.photos import Reading, Answer, to_record, read_photo, extract
from hepscape.cli import merge

ROOT = Path(__file__).resolve().parents[1]


@pytest.fixture
def data():
    return {
        "version": 1,
        "event": "Evento test",
        "records": [
            {
                "id": "TEST-01",
                "city": "Roma",
                "kit": "Roma",
                "source": "foto.jpg",
                "answers": {q: "Bella, bella" if q == "Q2" else 0 for q in QUESTIONS},
                "originals": {
                    q: "Bella, bella" if q == "Q2" else OPTIONS[q][0] for q in QUESTIONS
                },
                "reviewed": dict.fromkeys(QUESTIONS, True),
                "issues": {},
                "notes": {},
            }
        ],
    }


def edit_review(path, transform):
    with path.open(encoding="utf-8-sig", newline="") as h:
        reader = csv.DictReader(h)
        fields, rows = reader.fieldnames, list(reader)
    transform(rows)
    with path.open("w", encoding="utf-8-sig", newline="") as h:
        writer = csv.DictWriter(h, fieldnames=fields)
        writer.writeheader()
        writer.writerows(rows)


def test_coding_order():
    assert OPTIONS["Q8"] == OPTIONS["Q11"] == ["No", "Sì"]
    assert OPTIONS["Q9"][3] == "Preferisco non rispondere"
    assert len(OPTIONS["Q10"]) == 8
    assert coded("Q1", 0) == 0
    assert coded("Q2", "Curiosità") == "CURIOSITÀ"
    assert coded("Q2", None) == 999


def test_original_dataset_matches_every_excel_cell(tmp_path):
    data = load(ROOT / "data/ern2026/questionari_verificati.json", require_review=True)
    assert len(data["records"]) == 109
    target = tmp_path / "reproduced.xlsx"
    export_workbook(data, target)
    assert read_workbook(target) == read_workbook(
        ROOT / "data/ern2026/questionari.xlsx"
    )
    wb = load_workbook(target)
    assert wb["Dati codificati"].freeze_panes == "B7"
    assert len(wb["Dati codificati"].data_validations.dataValidation) == 10
    assert wb["Risposte originali"].max_row == 110


@pytest.mark.parametrize("value", [3, -1, "0", True, 999])
def test_invalid_category_rejected(data, value):
    data["records"][0]["answers"]["Q1"] = value
    with pytest.raises(ValueError):
        validate(data)


def test_pending_not_exportable(data, tmp_path):
    r = data["records"][0]
    r["reviewed"]["Q3"] = False
    r["answers"]["Q3"] = None
    r["issues"]["Q3"] = "Segno illeggibile"
    with pytest.raises(ValueError):
        export_workbook(data, tmp_path / "final.xlsx")
    export_workbook(data, tmp_path / "draft.xlsx", draft=True)
    assert (
        load_workbook(tmp_path / "draft.xlsx")["Dati codificati"]["D7"].value
        == "DA VERIFICARE"
    )
    with pytest.raises(ValueError, match="bozza"):
        read_workbook(tmp_path / "draft.xlsx")


def test_review_explicit_blank_and_code(data, tmp_path):
    r = data["records"][0]
    r["reviewed"]["Q3"] = False
    r["answers"]["Q3"] = None
    r["issues"]["Q3"] = "Segno illeggibile"
    path = tmp_path / "review.csv"
    export_review(data, path)
    edit_review(path, lambda rows: rows[2].update(confermato="SI"))
    with pytest.raises(ValueError, match="incerta"):
        apply_review(data, path)
    edit_review(path, lambda rows: rows[2].update(correzione="__BLANK__"))
    corrected = apply_review(data, path)
    assert corrected["records"][0]["answers"]["Q3"] is None
    assert corrected["records"][0]["reviewed"]["Q3"] is True
    edit_review(path, lambda rows: rows[2].update(correzione="2"))
    assert apply_review(data, path)["records"][0]["answers"]["Q3"] == 2
    assert data["records"][0]["answers"]["Q3"] is None


def test_stale_or_incomplete_review_rejected(data, tmp_path):
    path = tmp_path / "review.csv"
    export_review(data, path)
    changed = copy.deepcopy(data)
    changed["records"][0]["answers"]["Q1"] = 1
    with pytest.raises(ValueError, match="obsoleta"):
        apply_review(changed, path)
    edit_review(path, lambda rows: rows.pop())
    with pytest.raises(ValueError, match="incompleto"):
        apply_review(data, path)


def test_formula_like_text_preserved(data, tmp_path):
    data["records"][0]["answers"]["Q2"] = "=INTERESSANTE"
    path = tmp_path / "text.xlsx"
    export_workbook(data, path)
    c = load_workbook(path)["Dati codificati"]["C7"]
    assert c.value == "=INTERESSANTE" and c.data_type == "s"
    export_review(data, tmp_path / "review.csv")
    assert "'=INTERESSANTE" in (tmp_path / "review.csv").read_text(encoding="utf-8-sig")


def test_merge_multiple_cities_and_duplicate_detection(data):
    second = copy.deepcopy(data)
    second["records"][0].update(id="AVZ-01", city="Avezzano")
    combined = merge([data, second])
    assert [r["city"] for r in combined["records"]] == ["Roma", "Avezzano"]
    with pytest.raises(ValueError, match="duplicati"):
        merge([data, data])
    second["event"] = "altro"
    with pytest.raises(ValueError, match="eventi"):
        merge([data, second])


def reading():
    return Reading(
        questionnaire_visible=True,
        handwritten_id="42",
        answers=[
            Answer(
                question=q,
                status="selected",
                value="interessante" if q == "Q2" else "1",
                transcription="testo",
                note="",
            )
            for q in QUESTIONS
        ],
    )


def test_vision_mapping_and_review_gate():
    parsed = reading()
    parsed.answers[2] = Answer(
        question="Q3", status="blank", value=None, transcription="", note=""
    )
    parsed.answers[3] = Answer(
        question="Q4",
        status="uncertain",
        value=None,
        transcription="",
        note="due crocette",
    )
    record = to_record(
        parsed,
        identity="TEST-1",
        city="Roma",
        kit="Roma",
        source="test.jpg",
        digest="abc",
        model="test-model",
    )
    assert record["answers"]["Q8"] == 1
    assert record["answers"]["Q3"] is None
    assert record["issues"]["Q4"] == "due crocette"
    assert not any(record["reviewed"].values())


def test_duplicate_questions_rejected():
    parsed = reading()
    parsed.answers[-1] = parsed.answers[0]
    with pytest.raises(ValueError):
        to_record(
            parsed,
            identity="TEST-1",
            city="Roma",
            kit="Roma",
            source="test.jpg",
            digest="abc",
            model="test",
        )


def test_photo_request_contract_and_failure(tmp_path):
    path = tmp_path / "sample.jpg"
    Image.new("RGB", (400, 600), "white").save(path)
    calls = []

    def parse(**kwargs):
        calls.append(kwargs)
        return SimpleNamespace(status="completed", output_parsed=reading())

    client = SimpleNamespace(responses=SimpleNamespace(parse=parse))
    result = read_photo(client, path, "chosen-model")
    assert len(result.answers) == 11
    assert calls[0]["store"] is False
    assert calls[0]["model"] == "chosen-model"
    assert calls[0]["input"][1]["content"][1]["image_url"].startswith(
        "data:image/jpeg;base64,"
    )
    client.responses.parse = lambda **kw: SimpleNamespace(
        status="incomplete", output_parsed=None
    )
    with pytest.raises(ValueError, match="incompleta"):
        read_photo(client, path, "chosen-model")


def test_extract_identical_images_are_not_double_counted(tmp_path):
    first, second = tmp_path / "a.jpg", tmp_path / "b.jpg"
    Image.new("RGB", (100, 150), "white").save(first)
    second.write_bytes(first.read_bytes())
    client = SimpleNamespace(
        responses=SimpleNamespace(
            parse=lambda **kwargs: SimpleNamespace(
                status="completed", output_parsed=reading()
            )
        )
    )
    data, errors = extract(
        client,
        [first, second],
        tmp_path / "out",
        city="Roma",
        kit="Roma",
        event="test",
        prefix="T",
        model="test",
    )
    assert len(data["records"]) == 1 and len(errors) == 1
    assert (tmp_path / "out/bozza.json").exists()


def test_analysis_empty_q2_small_and_all_missing_groups(data, tmp_path):
    from hepscape.plots import generate
    from pypdf import PdfReader

    r = data["records"][0]
    r["answers"] = dict.fromkeys(QUESTIONS, None)
    second = copy.deepcopy(r)
    second.update(id="ALTRO-01", city="Avezzano", kit="Milano")
    data["records"].append(second)
    path = tmp_path / "empty.xlsx"
    export_workbook(data, path)
    generate(path, tmp_path / "report", "Test senza risposte", compare_kits=True)
    assert len(PdfReader(tmp_path / "report/HEPscape_raccolta_grafici.pdf").pages) == 18
    assert len(list((tmp_path / "report/grafici").glob("*.svg"))) == 17
    manifest = json.loads((tmp_path / "report/manifest.json").read_text())
    assert "15_citta_1" in manifest["figures"]
    assert "16_kit_1" in manifest["figures"]


def test_guided_output_prompts_and_safe_path(monkeypatch):
    from hepscape.cli import parser
    from hepscape.routing import destination

    monkeypatch.setattr("sys.stdin.isatty", lambda: True)
    responses = iter([" rOmA "])
    monkeypatch.setattr("builtins.input", lambda _: next(responses))
    args = parser().parse_args(["plots", "input.xlsx"])
    first, kit, city = destination(args)
    assert first.parent == Path("kits/roma/tutte-le-citta")
    assert kit == "Roma" and city is None
    args = parser().parse_args(
        ["plots", "input.xlsx", "--kit", "Roma", "--city", "../../Bari"]
    )
    path, _, _ = destination(args)
    assert path.parent == Path("kits/roma/bari")


def test_batch_routing_and_metadata_selection(monkeypatch):
    from hepscape.cli import parser
    from hepscape.routing import destination, select_rows

    monkeypatch.setattr("sys.stdin.isatty", lambda: False)
    with pytest.raises(ValueError, match="automatica"):
        destination(parser().parse_args(["plots", "input.xlsx"]))
    assert destination(
        parser().parse_args(["plots", "input.xlsx", "--out", "all"])
    ) == (Path("all"), None, None)
    with pytest.raises(ValueError, match="Kit ammessi"):
        destination(
            parser().parse_args(
                ["plots", "input.xlsx", "--kit", "Invalid", "--city", "Roma"]
            )
        )
    rows = [
        ("A",) + (0,) * 11 + ("Avezzano", "Roma"),
        ("B",) + (0,) * 11 + ("Bari", "Roma"),
        ("C",) + (0,) * 11 + ("Avezzano", "Pisa"),
    ]
    assert select_rows(rows, " roma ", "AVEZZANO") == rows[:1]
    with pytest.raises(ValueError, match="Nessuna scheda"):
        select_rows(rows, "Bari", "Avezzano")
    assert select_rows(rows) == rows
    assert select_rows(rows, kit="Roma") == rows[:2]
    assert select_rows(rows, city="Avezzano") == [rows[0], rows[2]]
    output, kit, city = destination(
        parser().parse_args(["plots", "input.xlsx", "--kit", "tutti"])
    )
    assert output.parent == Path("reports/tutti-i-kit/tutte-le-citta")
    assert kit is None and city is None


def test_filtered_report_manifest_and_no_false_output(data, tmp_path):
    from hepscape.plots import generate

    second = copy.deepcopy(data["records"][0])
    second.update(id="SECOND", city="Bari", kit="Bari")
    data["records"].append(second)
    source = tmp_path / "mixed.xlsx"
    export_workbook(data, source)
    generate(source, tmp_path / "selected", kit="Roma", city="Roma")
    report = json.loads((tmp_path / "selected/manifest.json").read_text())
    assert report["questionnaires"] == 1
    assert report["selection"] == {"kit": "Roma", "city": "Roma"}
    with pytest.raises(ValueError, match="Nessuna scheda"):
        generate(source, tmp_path / "absent", kit="Pisa", city="Roma")
    assert not (tmp_path / "absent").exists()


def test_all_kits_report_without_comparisons(data, tmp_path):
    from hepscape.plots import generate

    second = copy.deepcopy(data["records"][0])
    second.update(id="BARI-01", kit="Bari")
    data["records"].append(second)
    source = tmp_path / "all.xlsx"
    export_workbook(data, source)
    generate(source, tmp_path / "all")
    report = json.loads((tmp_path / "all/manifest.json").read_text())
    assert report["questionnaires"] == 2
    assert report["compare_kits"] is False
    assert not any(name.startswith("16_kit") for name in report["figures"])
    assert report["event"] == "Tutti gli eventi"
